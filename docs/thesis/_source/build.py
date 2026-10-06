# -*- coding: utf-8 -*-
import copy, sys, json, re
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import content as C

TEMPLATE = os.path.join(HERE, 'template.docx')
OUT = sys.argv[1]
PAGES = json.load(open(sys.argv[2])) if len(sys.argv) > 2 else {}

d = docx.Document(TEMPLATE)
body = d.element.body
els = list(body.iterchildren())
BLACK = RGBColor(0, 0, 0)

def set_text(p, text, color=BLACK, italic=None):
    runs = p.runs
    if not runs:
        r = p.add_run(text)
    else:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
        r = runs[0]
    r.font.color.rgb = color
    if italic is not None:
        r.font.italic = italic
    return r

def para_of(el):
    return docx.text.paragraph.Paragraph(el, d._body)

# ---------- cover ----------
set_text(para_of(els[6]), C.TITLE.upper())
set_text(para_of(els[10]), "Individual project")
tbl = docx.table.Table(els[14], d._body)
rows = tbl.rows
def fill_cell(cell, text):
    p = cell.paragraphs[0]
    set_text(p, text, italic=False)
fill_cell(rows[0].cells[1], "Obrijan Filip, FAF-232")
fill_cell(rows[0].cells[2], "")
fill_cell(rows[4].cells[0], "University Mentor")
fill_cell(rows[4].cells[1], "Ciutac Ștefănița")
fill_cell(rows[4].cells[2], "")
for i in (3, 2, 1):
    tr = rows[i]._tr
    tr.getparent().remove(tr)
set_text(para_of(els[16]), para_of(els[16]).text.replace("each student confirms their", "the student confirms their").replace("Students who did not contribute meaningfully are not permitted to sign.", "").strip())

# ---------- abstract ----------
set_text(para_of(els[29]), "ABSTRACT")
def placeholder_runs(p, text):
    r = set_text(p, text, color=RGBColor(0x59, 0x59, 0x59), italic=True)
    r.font.underline = False
    r.font.bold = False
    return r
placeholder_runs(para_of(els[32]), "[The abstract (300–500 words) will be written after Chapters 1–3 are completed.]")
placeholder_runs(para_of(els[33]), "Keywords: [5–7 keywords]")

sect_par = els[34]
# delete everything after the abstract section until the final sectPr
for el in els[35:]:
    if el.tag == qn('w:sectPr'):
        continue
    body.remove(el)

# ---------- helpers ----------
bm_id = [100]
TOC_ENTRIES = []   # (level, text, bookmark)

def add_bookmark(p, name):
    bs = OxmlElement('w:bookmarkStart'); bs.set(qn('w:id'), str(bm_id[0])); bs.set(qn('w:name'), name)
    be = OxmlElement('w:bookmarkEnd'); be.set(qn('w:id'), str(bm_id[0]))
    p._p.insert(1 if p._p.pPr is not None else 0, bs)
    p._p.append(be)
    bm_id[0] += 1

def tnr(r, size=None, bold=None):
    r.font.name = 'Times New Roman'
    rpr = r._r.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        rf.set(qn(a), 'Times New Roman')
    for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]
    if size: r.font.size = Pt(size)
    if bold is not None: r.font.bold = bold
    r.font.color.rgb = BLACK

def new_par(style=None, before=None):
    if before is not None:
        p = before.insert_paragraph_before('', style=style)
    else:
        p = d.add_paragraph(style=style)
    return p

def h1(text, before=None, toc=True):
    p = new_par('Heading 1', before)
    r = p.add_run(text); tnr(r, 13, True)
    if toc:
        name = '_Toc9%05d' % bm_id[0]
        add_bookmark(p, name); TOC_ENTRIES.append((1, text, name))
    return p

def h2(text, before=None, center=False, page_break=False, toc=True):
    p = new_par('Heading 2', before)
    pf = p.paragraph_format
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
    else:
        pf.first_line_indent = Cm(1.25)
    if page_break:
        pf.page_break_before = True
    r = p.add_run(text); tnr(r, 12, True)
    if toc:
        name = '_Toc9%05d' % bm_id[0]
        add_bookmark(p, name); TOC_ENTRIES.append((2, text, name))
    return p

def body_par(text, before=None, indent=True, highlight=False):
    p = new_par('Body Text 2026', before)
    if not indent:
        p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(text); tnr(r, 12)
    if highlight:
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p

def ph_par(text, before=None):
    p = new_par('Body Text 2026', before)
    r = p.add_run(text); tnr(r, 12); r.font.italic = True
    r.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
    return p

def dash_list(items, before=None):
    for i, it in enumerate(items):
        p = new_par('List Paragraph', before)
        p.paragraph_format.first_line_indent = None
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        numPr = p._p.get_or_add_pPr().get_or_add_numPr()
        numPr.get_or_add_ilvl().val = 0
        numPr.get_or_add_numId().val = 9
        r = p.add_run(it); tnr(r, 12)

def set_cell_shading(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def cell_text(cell, text, bold=False, highlight=False, center=False):
    lines = text.split('\n')
    p = cell.paragraphs[0]
    for k, line in enumerate(lines):
        if k > 0:
            p = cell.add_paragraph()
        p.style = d.styles['Normal']
        pf = p.paragraph_format
        pf.line_spacing = 1.0; pf.space_before = Pt(0); pf.space_after = Pt(0); pf.first_line_indent = Cm(0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(line); tnr(r, 11, bold)
        if highlight and line.startswith('['):
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW

def set_grid(t, widths):
    grid = t._tbl.tblGrid
    for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
        gc.set(qn('w:w'), str(int(w * 567)))
    for row in t.rows:
        for j, w in enumerate(widths):
            row.cells[j].width = Cm(w)

def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader'); e.set(qn('w:val'), 'true'); trPr.append(e)

def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit'); e.set(qn('w:val'), 'true'); trPr.append(e)

last_table = [None]

def table(caption, header, rows, widths, highlight=False, before=None):
    if caption is None and last_table[0] is not None:
        t = last_table[0]
        hr = t.add_row()
        for j, h in enumerate(header):
            cell_text(hr.cells[j], h, bold=True, center=True); set_cell_shading(hr.cells[j], 'D9D9D9')
        for row in rows:
            rr = t.add_row(); no_split(rr)
            for j, v in enumerate(row):
                cell_text(rr.cells[j], v, highlight=highlight)
        return t
    cp = new_par('Normal', before)
    cp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    cp.paragraph_format.keep_with_next = True
    cp.paragraph_format.space_before = Pt(6)
    r = cp.add_run(caption); tnr(r, 12, True)
    t = d.add_table(rows=1 + len(rows), cols=len(header))
    if before is not None:
        before._p.addprevious(t._tbl)
    t.style = d.styles['Table Grid']
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        cell_text(c, h, bold=True, center=True); set_cell_shading(c, 'D9D9D9')
    set_repeat_header(t.rows[0])
    for i, row in enumerate(rows):
        no_split(t.rows[i + 1])
        for j, v in enumerate(row):
            center = len(header) > 4 and j > 0
            cell_text(t.rows[i + 1].cells[j], v, highlight=highlight, center=center)
    set_grid(t, widths)
    # keep caption with table: keep_with_next on all paragraphs of first row
    for c in t.rows[0].cells:
        for p in c.paragraphs:
            p.paragraph_format.keep_with_next = True
    last_table[0] = t
    # small spacer after the table
    sp = new_par('Normal', before)
    sp.paragraph_format.space_after = Pt(0)
    return t

def figure(path, caption, width_cm, before=None):
    p = new_par('Normal', before)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_picture(path, width=Cm(width_cm))
    c = new_par('Normal', before)
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    c.paragraph_format.space_after = Pt(6)
    r = c.add_run(caption); tnr(r, 12, True)

def render(blocks):
    for b in blocks:
        kind = b[0]
        if kind == 'h1': h1(b[1])
        elif kind == 'h2': h2(b[1])
        elif kind == 'p': body_par(b[1])
        elif kind == 'pc': body_par(b[1], indent=False)
        elif kind == 'ph': body_par(b[1], highlight=True)
        elif kind == 'list': dash_list(b[1])
        elif kind == 'fig': figure(b[1], b[2], b[3])
        elif kind == 'table':
            table(b[1], b[2], b[3], b[4], highlight=(len(b) > 5 and b[5] == 'highlight'))

# ---------- list of abbreviations (inside the abstract section, before Introduction) ----------
sp = para_of(sect_par)
h1("LIST OF ABBREVIATIONS", before=sp)
t = d.add_table(rows=len(C.ABBREVIATIONS), cols=2)
sp._p.addprevious(t._tbl)
for i, (a, desc) in enumerate(C.ABBREVIATIONS):
    cell_text(t.rows[i].cells[0], a, bold=True)
    cell_text(t.rows[i].cells[1], "– " + desc)
    t.rows[i].cells[0].width = Cm(3); t.rows[i].cells[1].width = Cm(14)
    for c in t.rows[i].cells:
        for p in c.paragraphs:
            p.paragraph_format.space_after = Pt(4)
            for r in p.runs: r.font.size = Pt(12)
set_grid(t, [3, 14])

# ---------- main body ----------
h1("INTRODUCTION")
ph_par("[The introduction (1–2 pages) will be written after the chapters are completed: relevance of the topic, "
       "the problem, the aim, the objectives, the object of research, the methods used and the structure of the paper.]")

render(C.CH1)

h1("2 UML MODELLING OF THE SYSTEM")
ph_par("[Chapter 2 will be completed at the next stage: use case model, activity diagrams, sequence and state "
       "diagrams, class diagram, component and package diagrams, architecture and deployment.]")
h1("3 IMPLEMENTATION OF THE SYSTEM")
ph_par("[Chapter 3 will be completed at the implementation stage of the bachelor's project.]")
h1("CONCLUSIONS")
ph_par("[The general conclusions (2–3 pages) will be written at the end of the work.]")

h1("REFERENCES")
for i, ref in enumerate(C.REFERENCES, 1):
    p = new_par('Normal')
    pf = p.paragraph_format
    pf.left_indent = Cm(1.0); pf.first_line_indent = Cm(-1.0)
    pf.space_after = Pt(3)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("[%d]\t%s" % (i, ref)); tnr(r, 12)
    pf.tab_stops.add_tab_stop(Cm(1.0))

h1("APPENDICES")
h2("APPENDIX A – Project proposal", center=True)
p = body_par("The project proposal approved for the bachelor's thesis is reproduced below in English.")
t = d.add_table(rows=len(C.PROPOSAL), cols=2)
t.style = d.styles['Table Grid']; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (lab, lines) in enumerate(C.PROPOSAL):
    cell_text(t.rows[i].cells[0], lab, bold=True)
    c = t.rows[i].cells[1]
    if len(lines) == 1:
        cell_text(c, lines[0])
    else:
        cell_text(c, "\n".join("- " + l for l in lines))
    t.rows[i].cells[0].width = Cm(3.5); t.rows[i].cells[1].width = Cm(13.5)
    for cc in t.rows[i].cells:
        for pp in cc.paragraphs:
            pp.paragraph_format.space_after = Pt(3)
            for r in pp.runs: r.font.size = Pt(11)
set_grid(t, [3.5, 13.5])
d.add_paragraph()
p = body_par("Student: Obrijan Filip, group FAF-232", indent=False)
p = body_par("Scientific supervisor: Ciutac Ștefănița", indent=False)

h2("APPENDIX B – Questionnaire and interview guide", center=True, page_break=True)
body_par(C.QUESTIONNAIRE_INTRO)
p = body_par("Questions for producers (sellers):", indent=False); p.runs[0].font.bold = True
for i, q in enumerate(C.Q_PRODUCERS):
    qp = body_par("%s) %s" % ("abcdefghij"[i], q)); qp.paragraph_format.first_line_indent = Cm(-0.6); qp.paragraph_format.left_indent = Cm(1.85)
p = body_par("Questions for wholesale buyers (distributors, HoReCa, retail):", indent=False); p.runs[0].font.bold = True
for i, q in enumerate(C.Q_BUYERS):
    qp = body_par("%s) %s" % ("abcdefghij"[i], q)); qp.paragraph_format.first_line_indent = Cm(-0.6); qp.paragraph_format.left_indent = Cm(1.85)

# ---------- pagination rules: no table/figure at the top or bottom of a page, never split ----------
def kwn(el):
    if el is not None and el.tag == qn('w:p'):
        docx.text.paragraph.Paragraph(el, d._body).paragraph_format.keep_with_next = True

def prev_text_par(el):
    p = el.getprevious()
    while p is not None and p.tag == qn('w:p') and not ''.join(t.text or '' for t in p.iter(qn('w:t'))).strip():
        p = p.getprevious()
    return p

started = False
for el in list(body.iterchildren()):
    if el.tag == qn('w:p') and ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip() == 'INTRODUCTION':
        started = True
    if not started:
        continue
    if el.tag == qn('w:tbl'):
        # whole table on one page and kept together with the text that follows it
        for p in el.iter(qn('w:p')):
            kwn(p)
        nxt = el.getnext()
        while nxt is not None and nxt.tag == qn('w:p') and not ''.join(t.text or '' for t in nxt.iter(qn('w:t'))).strip():
            kwn(nxt); nxt = nxt.getnext()
        # caption and at least the end of the previous paragraph stay above the table
        cap = el.getprevious()
        kwn(cap)
        before_cap = prev_text_par(cap)
        # do not glue a paragraph that already follows another table/figure, otherwise long chains form
        bb = before_cap.getprevious() if before_cap is not None else None
        while bb is not None and bb.tag == qn('w:p') and not ''.join(t.text or '' for t in bb.iter(qn('w:t'))).strip() and bb.find('.//' + qn('w:drawing')) is None:
            bb = bb.getprevious()
        if bb is None or bb.tag != qn('w:tbl'):
            kwn(before_cap)
    elif el.tag == qn('w:p') and el.find('.//' + qn('w:drawing')) is not None:
        kwn(el)                       # picture stays with its caption
        kwn(el.getnext())             # caption stays with the following text
        kwn(prev_text_par(el))        # previous text stays above the picture

# ---------- table of contents ----------
sdt = body.find(qn('w:sdt'))
sc = sdt.find(qn('w:sdtContent'))
ps = sc.findall(qn('w:p'))
first, end_p, sect_p = ps[0], ps[-2], ps[-1]
for p_ in ps[:-2]:
    sc.remove(p_)

def mk_run(text=None, fld=None, instr=None, tab=False):
    r = OxmlElement('w:r')
    rpr = OxmlElement('w:rPr'); nf = OxmlElement('w:noProof'); rpr.append(nf)
    rf = OxmlElement('w:rFonts')
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'): rf.set(qn(a), 'Times New Roman')
    rpr.insert(0, rf)
    r.append(rpr)
    if fld:
        f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), fld); r.append(f)
    if instr:
        it = OxmlElement('w:instrText'); it.set(qn('xml:space'), 'preserve'); it.text = instr; r.append(it)
    if tab:
        r.append(OxmlElement('w:tab'))
    if text is not None:
        t_ = OxmlElement('w:t'); t_.set(qn('xml:space'), 'preserve'); t_.text = text; r.append(t_)
    return r

entries = [(1, "ABSTRACT", None), (1, "LIST OF ABBREVIATIONS", None)] + TOC_ENTRIES
# abstract heading bookmark
abs_p = para_of(els[29])
add_bookmark(abs_p, '_Toc900001')
# abbreviations heading bookmark is TOC_ENTRIES[0]
entries = [(1, "ABSTRACT", '_Toc900001')] + TOC_ENTRIES
for k, (lvl, text, bm) in enumerate(entries):
    p = OxmlElement('w:p')
    ppr = OxmlElement('w:pPr'); st = OxmlElement('w:pStyle'); st.set(qn('w:val'), 'TOC%d' % lvl); ppr.append(st)
    p.append(ppr)
    if k == 0:
        p.append(mk_run(fld='begin')); p.append(mk_run(instr=' TOC \\o "1-3" \\h \\z \\u ')); p.append(mk_run(fld='separate'))
    hl = OxmlElement('w:hyperlink'); hl.set(qn('w:anchor'), bm); hl.set(qn('w:history'), '1')
    hl.append(mk_run(text=text))
    hl.append(mk_run(tab=True))
    hl.append(mk_run(fld='begin')); hl.append(mk_run(instr=' PAGEREF %s \\h ' % bm)); hl.append(mk_run(fld='separate'))
    hl.append(mk_run(text=str(PAGES.get(text, ''))))
    hl.append(mk_run(fld='end'))
    p.append(hl)
    end_p.addprevious(p)

d.save(OUT)
json.dump([e[1] for e in entries], open(OUT + '.headings.json', 'w'))
print('saved', OUT)
